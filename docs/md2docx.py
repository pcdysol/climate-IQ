#!/usr/bin/env python
"""Minimal Markdown -> DOCX converter tailored to the ClimateIQ architecture report.

Handles: # title, ##/###/#### headings, --- page breaks, pipe tables,
``` code/diagram blocks, - bullet lists (one level of nesting), > note callouts,
and inline **bold** / *italic* / `code`.
"""
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

SRC = sys.argv[1] if len(sys.argv) > 1 else "ClimateIQ_Architecture_Report.md"
OUT = sys.argv[2] if len(sys.argv) > 2 else "ClimateIQ_Architecture_Report.docx"

ACCENT = RGBColor(0x1F, 0x4E, 0x79)   # dark blue (headings)
ACCENT2 = RGBColor(0x2E, 0x74, 0xB5)  # lighter blue (sub headings)
SHADE = "F2F2F2"

doc = Document()

# Base style
normal = doc.styles["Normal"]
normal.font.name = "Calibri"
normal.font.size = Pt(10.5)

def shade_cell(cell, color):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), color)
    tcPr.append(shd)

def add_runs(paragraph, text):
    """Parse inline **bold**, *italic*, `code` into runs."""
    token = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")
    for part in token.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2]); r.bold = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1])
            r.font.name = "Consolas"; r.font.size = Pt(9.5)
        elif part.startswith("*") and part.endswith("*"):
            r = paragraph.add_run(part[1:-1]); r.italic = True
        else:
            paragraph.add_run(part)

def add_code_block(lines):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.left_indent = Inches(0.1)
    pf.space_before = Pt(4); pf.space_after = Pt(8)
    # light shading via paragraph border/background
    pPr = p._p.get_or_add_pPr()
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), SHADE)
    pPr.append(shd)
    run = p.add_run("\n".join(lines))
    run.font.name = "Consolas"; run.font.size = Pt(8.5)

def add_table(rows):
    header, body = rows[0], rows[1:]
    t = doc.add_table(rows=1, cols=len(header))
    t.style = "Table Grid"
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    for i, cell_text in enumerate(header):
        c = t.rows[0].cells[i]
        c.paragraphs[0].text = ""
        add_runs(c.paragraphs[0], cell_text)
        for r in c.paragraphs[0].runs:
            r.bold = True; r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
        shade_cell(c, "1F4E79")
    for row in body:
        cells = t.add_row().cells
        for i in range(len(header)):
            txt = row[i] if i < len(row) else ""
            cells[i].paragraphs[0].text = ""
            add_runs(cells[i].paragraphs[0], txt)
    doc.add_paragraph()

def parse_table_row(line):
    return [c.strip() for c in line.strip().strip("|").split("|")]

with open(SRC, encoding="utf-8") as f:
    lines = f.read().split("\n")

i = 0
n = len(lines)
while i < n:
    line = lines[i]
    stripped = line.strip()

    # Code / diagram block
    if stripped.startswith("```"):
        block = []
        i += 1
        while i < n and not lines[i].strip().startswith("```"):
            block.append(lines[i]); i += 1
        add_code_block(block)
        i += 1
        continue

    # Table
    if stripped.startswith("|") and i + 1 < n and re.match(r"^\|[\s:|-]+\|?$", lines[i+1].strip()):
        rows = [parse_table_row(stripped)]
        i += 2  # skip header + separator
        while i < n and lines[i].strip().startswith("|"):
            rows.append(parse_table_row(lines[i].strip())); i += 1
        add_table(rows)
        continue

    # Horizontal rule -> page break
    if stripped == "---":
        doc.add_page_break()
        i += 1
        continue

    # Headings
    if stripped.startswith("#"):
        m = re.match(r"^(#+)\s+(.*)$", stripped)
        level, text = len(m.group(1)), m.group(2)
        text = re.sub(r"\*\*(.+?)\*\*", r"\1", text)
        text = text.replace("`", "")
        if level == 1:
            p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r = p.add_run(text); r.bold = True; r.font.size = Pt(26)
            r.font.color.rgb = ACCENT
        elif level == 2:
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(10)
            r = p.add_run(text); r.bold = True; r.font.size = Pt(17)
            r.font.color.rgb = ACCENT
            # underline-ish bottom border
            pPr = p._p.get_or_add_pPr()
            pbdr = OxmlElement("w:pBdr"); bottom = OxmlElement("w:bottom")
            bottom.set(qn("w:val"), "single"); bottom.set(qn("w:sz"), "6")
            bottom.set(qn("w:space"), "2"); bottom.set(qn("w:color"), "1F4E79")
            pbdr.append(bottom); pPr.append(pbdr)
        elif level == 3:
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(8)
            r = p.add_run(text); r.bold = True; r.font.size = Pt(13.5)
            r.font.color.rgb = ACCENT2
        else:
            p = doc.add_paragraph(); p.paragraph_format.space_before = Pt(6)
            r = p.add_run(text); r.bold = True; r.font.size = Pt(11.5)
            r.font.color.rgb = ACCENT2
        i += 1
        continue

    # Blockquote / note
    if stripped.startswith(">"):
        block = []
        while i < n and lines[i].strip().startswith(">"):
            block.append(lines[i].strip().lstrip(">").strip()); i += 1
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.25)
        pPr = p._p.get_or_add_pPr()
        shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:fill"), "FFF8E1")
        pPr.append(shd)
        add_runs(p, " ".join(block))
        i += 1
        continue

    # Bullet list
    if re.match(r"^(\s*)[-*]\s+", line):
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        indent = len(m.group(1))
        style = "List Bullet 2" if indent >= 2 else "List Bullet"
        p = doc.add_paragraph(style=style)
        add_runs(p, m.group(2))
        i += 1
        continue

    # Blank line
    if stripped == "":
        i += 1
        continue

    # Normal paragraph
    p = doc.add_paragraph()
    add_runs(p, stripped)
    i += 1

doc.save(OUT)
print("Wrote", OUT)
